from pathlib import Path
import argparse
import json
import re
from datetime import datetime
from docx.styles.style import _ParagraphStyle

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


def load_json(path: str):
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def safe_slug(text: str):
    text = text.lower().strip()
    text = re.sub(r"[^a-z0-9]+", "-", text)
    return text.strip("-")[:100]


def add_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(20)


def add_heading(doc, text, level=1):
    doc.add_heading(text, level=level)


def add_paragraph(doc, text):
    if text:
        doc.add_paragraph(str(text))


def add_bullet(doc, text):
    if text:
        doc.add_paragraph(str(text), style="List Bullet")


def add_key_value_table(doc, data: dict):
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"

    for key, value in data.items():
        row = table.add_row().cells
        row[0].text = str(key)
        row[1].text = str(value or "")

    doc.add_paragraph()


def add_strengths_table(doc, items):
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "JD Requirement"
    table.rows[0].cells[1].text = "CV Evidence"

    for item in items:
        row = table.add_row().cells
        row[0].text = item.get("requirement", "")
        row[1].text = item.get("evidence", "")
        
    doc.add_paragraph()


def add_weakness_table(doc, items):
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"

    headers = ["Gap", "Severity", "Adjacent Experience", "Mitigation"]
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h

    for item in items:
        row = table.add_row().cells
        row[0].text = item.get("gap", "")
        row[1].text = item.get("severity", "")
        row[2].text = item.get("adjacent_experience", "")
        row[3].text = item.get("mitigation", "")

    doc.add_paragraph()


def add_keyword_table(doc, items):
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"

    headers = ["Keyword", "Status", "Explanation"]
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h

    for item in items:
        row = table.add_row().cells
        row[0].text = item.get("keyword", "")
        row[1].text = item.get("status", "")
        row[2].text = item.get("explanation", "")

    doc.add_paragraph()


def add_cv_bullets_table(doc, items):
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"

    table.rows[0].cells[0].text = "CV Section"
    table.rows[0].cells[1].text = "Suggested Ready-to-Paste Content"

    for item in items:
        row = table.add_row().cells
        row[0].text = item.get("section", "")
        row[1].text = item.get("recommended_bullet", item.get("bullet", ""))

    doc.add_paragraph()


def build_docx(input_path: str):
    data = load_json(input_path)

    app = data.get("application_decision", {})
    company = app.get("company", "Unknown Company")
    role = app.get("role", "Unknown Role")

    doc = Document()

    section = doc.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)

    normal_style = doc.styles["Normal"]

    if isinstance(normal_style, _ParagraphStyle):
        normal_style.font.name = "Calibri"
        normal_style.font.size = Pt(10)

    add_title(doc, "Job-Specific CV Optimization Report")
    add_paragraph(doc, f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    doc.add_paragraph()

    add_heading(doc, "1. Application Assessment")
    add_key_value_table(doc, {
        "Company": company,
        "Role": role,
        "ATS Match Score": app.get("score"),
        "Recommendation": app.get("recommendation"),
        "Why Apply": app.get("why_apply"),
        "Main Gap": app.get("main_gap"),
    })

    add_heading(doc, "2. Strong Points")
    add_paragraph(doc, "Existing strengths already aligned to the role requirements.")
    add_strengths_table(doc, data.get("strong_points", []))

    add_heading(doc, "3. Weak Points / Gaps")
    add_paragraph(doc, "Gaps identified from the Phase 1 evaluation report and matched with adjacent CV evidence.")
    add_weakness_table(doc, data.get("weak_points", []))

    add_heading(doc, "4. ATS Keyword Analysis")
    add_keyword_table(doc, data.get("keyword_analysis", []))

    add_heading(doc, "5. Optimized Professional Summary")
    summary = data.get("professional_summary", {})
    add_heading(doc, "Current Summary", level=2)
    add_paragraph(doc, summary.get("current", ""))
    add_heading(doc, "Ready-to-Paste Optimized Summary", level=2)
    add_paragraph(doc, summary.get("optimized", ""))

    add_heading(doc, "6. ATS-Ready CV Improvements")
    add_cv_bullets_table(doc, data.get("cv_bullets_to_add", []))

    add_heading(doc, "7. LinkedIn Optimization")
    linkedin = data.get("linkedin_optimization", {})
    add_paragraph(doc, f"Suggested Headline: {linkedin.get('headline', '')}")
    for note in linkedin.get("notes", []):
        add_bullet(doc, note)

    add_heading(doc, "8. Interview Preparation")
    interview = data.get("interview_preparation", {})

    add_heading(doc, "STAR Stories to Prepare", level=2)
    for story in interview.get("structured_star_stories", []):
        if "raw" in story:
            add_bullet(doc, story.get("raw"))
        else:
            add_bullet(doc, f"Requirement: {story.get('requirement', '')}")
            add_paragraph(doc, f"Situation: {story.get('situation', '')}")
            add_paragraph(doc, f"Task: {story.get('task', '')}")
            add_paragraph(doc, f"Action: {story.get('action', '')}")
            add_paragraph(doc, f"Result: {story.get('result', '')}")

    add_heading(doc, "Red-Flag Questions", level=2)
    for item in interview.get("red_flags", []):
        add_bullet(doc, f"{item.get('question', '')} — {item.get('recommended_response', '')}")

    add_heading(doc, "9. Final Recommendation")
    final = data.get("final_recommendation", {})
    add_key_value_table(doc, {
        "Decision": final.get("decision"),
        "Confidence": final.get("confidence"),
        "Rationale": final.get("rationale"),
        "Main Gap": final.get("main_gap"),
        "Recommended Action": final.get("action"),
    })

    add_heading(doc, "10. Guardrail Note")
    add_paragraph(
        doc,
        "This report is an optimization advisor. It does not automatically rewrite the CV. "
        "All recommendations must be reviewed by the candidate before applying."
    )

    output_dir = Path("data/cv_optimization/reports")
    output_dir.mkdir(parents=True, exist_ok=True)

    output_path = output_dir / f"{safe_slug(company)}-cv-optimization-report.docx"
    doc.save(str(output_path))

    print()
    print("DOCX report created successfully")
    print(output_path)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    args = parser.parse_args()

    build_docx(args.input)


if __name__ == "__main__":
    main()