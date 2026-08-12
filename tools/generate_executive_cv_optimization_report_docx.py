from pathlib import Path
import argparse
import json
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt, RGBColor


REPORT_DIR = Path("data/cv_optimization/reports")
TOP_LEVEL_SECTION_ALLOWLIST = {
    "A. Job Fit Summary",
    "B. Missing Keywords from JD",
    "C. JD Requirements Match",
    "D. Gaps & Mitigation",
    "E. CV Changes",
}


def load_json(path: str):
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def safe_slug(text: str):
    text = str(text or "unknown").lower().strip()
    text = re.sub(r"[^a-z0-9]+", "-", text)
    return text.strip("-")[:100] or "unknown"


def clean_text(text):
    text = str(text or "")
    replacements = {
        "\u2010": "-",
        "\u2011": "-",
        "\u2012": "-",
        "\u2013": "-",
        "\u2014": "-",
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
    }

    for source, replacement in replacements.items():
        text = text.replace(source, replacement)

    return re.sub(r"\s+", " ", text).strip()


def set_style_font(style, name="Calibri", size=9.5, color=None, bold=False):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold

    if color:
        style.font.color.rgb = RGBColor.from_string(color)


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    normal = doc.styles["Normal"]
    set_style_font(normal, size=9.5)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing = 1.05

    for style_name, size, color, before, after in [
        ("Heading 1", 11.5, "2E74B5", 6, 2),
        ("Heading 2", 10.2, "1F4D78", 3, 1),
    ]:
        style = doc.styles[style_name]
        set_style_font(style, size=size, color=color, bold=True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    bullet_style = doc.styles["List Bullet"]
    set_style_font(bullet_style, size=9.2)
    bullet_style.paragraph_format.left_indent = Inches(0.23)
    bullet_style.paragraph_format.first_line_indent = Inches(-0.12)
    bullet_style.paragraph_format.space_after = Pt(1.5)
    bullet_style.paragraph_format.line_spacing = 1.03


def add_title(doc, company, role):
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(1)
    run = title.add_run("Executive CV Optimization Report")
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(15.5)
    run.font.color.rgb = RGBColor.from_string("0B2545")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(5)
    run = subtitle.add_run(f"{clean_text(company)} | {clean_text(role)}")
    run.font.name = "Calibri"
    run.font.size = Pt(8.8)
    run.font.color.rgb = RGBColor.from_string("555555")


def add_heading(doc, text, level=1):
    doc.add_heading(clean_text(text), level=level)


def add_top_level_heading(doc, text):
    text = clean_text(text)
    if text not in TOP_LEVEL_SECTION_ALLOWLIST:
        return
    add_heading(doc, text)


def add_label_value(doc, label, value):
    value = clean_text(value)

    if not value:
        return

    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(1.5)
    label_run = paragraph.add_run(f"{label}: ")
    label_run.bold = True
    label_run.font.name = "Calibri"
    value_run = paragraph.add_run(value)
    value_run.font.name = "Calibri"

def add_inline_keywords(doc, keywords):
    text = " | ".join(clean_text(keyword) for keyword in keywords)

    if not text:
        return

    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(9.5)


def add_bullet(doc, text):
    text = clean_text(text)

    if not text:
        return

    paragraph = doc.add_paragraph(text, style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(1.5)


def set_table_column_widths(table, widths_cm):
    if not widths_cm:
        return

    table.autofit = False
    for row in table.rows:
        for index, width_cm in enumerate(widths_cm):
            if index < len(row.cells) and width_cm is not None:
                row.cells[index].width = Cm(width_cm)


def add_table(doc, headers, rows, widths_cm=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"

    for index, header in enumerate(headers):
        run = table.rows[0].cells[index].paragraphs[0].add_run(clean_text(header))
        run.bold = True
        run.font.name = "Calibri"
        run.font.size = Pt(8.5)

    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            run = cells[index].paragraphs[0].add_run(clean_text(value))
            run.font.name = "Calibri"
            run.font.size = Pt(8.2)

    set_table_column_widths(table, widths_cm)


def add_jd_requirements_match(doc, items):
    rows = [
        [item.get("jd_requirement"), item.get("status") or item.get("match")]
        for item in items
        if item.get("jd_requirement")
    ]
    add_table(doc, ["JD Requirement", "Status"], rows, [16.7, 2.18])


def add_gaps_and_mitigation(doc, items):
    rows = []
    for item in items:
        if not item.get("gap") or not item.get("mitigation"):
            continue

        severity = clean_text(item.get("severity"))
        if re.match(r"^hard\b", severity, flags=re.IGNORECASE):
            severity = "Major"
        else:
            severity = "Minor"

        rows.append([item.get("gap"), severity, item.get("mitigation")])

    add_table(doc, ["Gap", "Severity", "Mitigation"], rows, [5.7, 2.18, 11.0])


def add_cv_changes(doc, items):
    rows = []
    for item in items:
        section = clean_text(item.get("section"))
        if not section or "skill" in section.lower() or "certification" in section.lower():
            continue

        rows.append([
            item.get("number"),
            section,
            item.get("current"),
            item.get("proposed_change"),
            item.get("why"),
        ])

    if not rows:
        return False

    add_table(
        doc,
        ["#", "Section", "Current state", "Proposed change", "Why"],
        rows,
        [0.7, 2.5, 3.6, 5.0, 7.08],
    )
    return True


def add_summary_change(doc, item):
    summary = clean_text(item.get("recommended_summary"))

    if not summary:
        return

    add_heading(doc, "Professional Summary", level=2)
    paragraph = doc.add_paragraph(summary)
    paragraph.paragraph_format.space_after = Pt(2)


def add_experience_change(doc, item):
    section = clean_text(item.get("section"))
    bullets = item.get("recommended_bullets", [])

    if not section or not bullets:
        return

    add_heading(doc, section, level=2)

    for bullet in bullets:
        add_bullet(doc, bullet)


def add_required_cv_changes(doc, items):
    rendered = False
    for item in items:
        if item.get("recommended_summary"):
            add_summary_change(doc, item)
            rendered = True
            continue
        if item.get("recommended_bullets"):
            add_experience_change(doc, item)
            rendered = True
    return rendered


def add_missing_skills(doc, missing_skills):
    category_labels = {
        "functional": "Functional",
        "technical": "Technical",
        "domain": "Domain",
        "certifications": "Certifications",
    }

    for category in ["functional", "technical", "domain", "certifications"]:
        skills = [
            clean_text(skill)
            for skill in missing_skills.get(category, [])
            if clean_text(skill)
        ]

        if not skills:
            continue

        add_heading(doc, category_labels[category], level=2)

        for skill in skills:
            add_bullet(doc, skill)


def add_interview_risks(doc, risks):
    for item in risks:
        risk = clean_text(item.get("risk"))
        positioning = clean_text(item.get("positioning"))

        if not risk and not positioning:
            continue

        if risk:
            risk_paragraph = doc.add_paragraph()
            risk_paragraph.paragraph_format.space_after = Pt(0.5)
            risk_run = risk_paragraph.add_run("Risk: ")
            risk_run.bold = True
            risk_run.font.name = "Calibri"
            value_run = risk_paragraph.add_run(risk)
            value_run.font.name = "Calibri"

        if positioning:
            positioning_paragraph = doc.add_paragraph()
            positioning_paragraph.paragraph_format.space_after = Pt(2)
            label_run = positioning_paragraph.add_run("Positioning: ")
            label_run.bold = True
            label_run.font.name = "Calibri"
            positioning_run = positioning_paragraph.add_run(positioning)
            positioning_run.font.name = "Calibri"


def build_docx(input_path: str, output_path: str | Path | None = None):
    draft = load_json(input_path)
    executive_report = draft.get("executive_report", {})

    if not executive_report:
        raise ValueError("Input draft does not contain executive_report.")

    summary = executive_report.get("job_fit_summary", {})
    company = summary.get("company") or "Unknown Company"
    role = summary.get("role") or "Unknown Role"

    doc = Document()
    configure_document(doc)

    add_title(doc, company, role)

    add_top_level_heading(doc, "A. Job Fit Summary")
    add_label_value(doc, "Company", company)
    add_label_value(doc, "Role", role)
    add_label_value(doc, "Score", summary.get("score"))
    add_label_value(doc, "Recommendation", summary.get("recommendation"))
    add_label_value(doc, "Fit Type", summary.get("fit_type"))
    add_label_value(doc, "Seniority Match", summary.get("seniority_match"))
    add_label_value(doc, "Why Apply", summary.get("why_apply"))
    add_label_value(doc, "Main Gap", summary.get("main_gap"))

    missing_keywords = [
        clean_text(keyword)
        for keyword in executive_report.get("missing_keywords", [])
        if clean_text(keyword)
    ]

    if missing_keywords:
        add_top_level_heading(doc, "B. Missing Keywords from JD")
        add_inline_keywords(doc, missing_keywords)

    jd_requirements_match = executive_report.get("jd_requirements_match", [])
    add_top_level_heading(doc, "C. JD Requirements Match")
    add_jd_requirements_match(doc, jd_requirements_match)

    gaps_and_mitigation = executive_report.get("gaps_and_mitigation", [])
    add_top_level_heading(doc, "D. Gaps & Mitigation")
    add_gaps_and_mitigation(doc, gaps_and_mitigation)

    cv_changes = executive_report.get("cv_changes", [])
    required_cv_changes = executive_report.get("required_cv_changes", [])
    add_top_level_heading(doc, "E. CV Changes")
    if not add_cv_changes(doc, cv_changes):
        add_required_cv_changes(doc, required_cv_changes)

    if output_path is None:
        output_path = REPORT_DIR / f"{safe_slug(company)}-{safe_slug(role)}-executive-report.docx"
    else:
        output_path = Path(output_path)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))

    print(output_path)

    return output_path


def main():
    parser = argparse.ArgumentParser(
        description="Generate one-page Executive CV Optimization Report DOCX."
    )
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", help="Optional output DOCX path.")
    args = parser.parse_args()

    build_docx(args.input, args.output)


if __name__ == "__main__":
    main()
